import io
import json
import sqlite3
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from modules.config import ChunkConfig
from modules.database_manager import DatabaseManager
from modules.documents import corpus_fingerprint, extract_document, load_knowledge
from modules.errors import DocumentError
from modules.text_processing import chunk_documents


def text_pdf(texts):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        content = DecodedStreamObject()
        content.set_data(f'BT /F1 12 Tf 40 700 Td ({text}) Tj ET'.encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(content)
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def test_import_preserves_fields_answers_sources_and_is_idempotent(tmp_path):
    db = DatabaseManager(tmp_path / 'nav.db')
    db.initialize()
    path = Path('database/knowledge_base.json')
    db.import_demo_documents(path)
    before = db.get_documents()
    db.import_demo_documents(path)
    assert db.get_documents() == before
    assert len(before) == 32
    for original, imported in zip(load_knowledge(path), before, strict=False):
        assert all(imported['metadata'][k] == v for k, v in original.items())
        text = imported['pages'][0]['text']
        assert original['answer'] in text
        assert all(step in text for step in original['steps'])
        assert json.loads(db.get_original(imported['id'])) == original


def test_json_list_and_failed_import_does_not_destroy_previous_data(tmp_path):
    items = load_knowledge('database/knowledge_base.json')
    path = tmp_path / 'items.json'
    path.write_text(json.dumps(items), encoding='utf-8')
    db = DatabaseManager(tmp_path / 'db', path)
    db.initialize()
    path.write_text(json.dumps([items[0], items[0]]), encoding='utf-8')
    with pytest.raises(DocumentError, match='уникальные'):
        db.import_knowledge_base()
    assert len(db.get_knowledge_items()) == 32
    with pytest.raises(DocumentError):
        db.import_demo_documents(path)
    assert db.get_documents() == []


def test_connections_close_on_success_and_rollback_on_failure(tmp_path):
    db = DatabaseManager(tmp_path / 'db')
    db.initialize()
    with db._connect() as connection:
        connection.execute('SELECT 1')
    with pytest.raises(sqlite3.ProgrammingError, match='closed'):
        connection.execute('SELECT 1')
    with pytest.raises(RuntimeError):
        with db._connect() as failed:
            failed.execute("INSERT INTO history(query,score,created_at) VALUES ('test',1,'now')")
            raise RuntimeError('rollback')
    with pytest.raises(sqlite3.ProgrammingError, match='closed'):
        failed.execute('SELECT 1')
    assert db.get_history() == []
    # On Windows a lingering SQLite handle prevents this rename/delete scenario.
    moved = tmp_path / 'renamed.db'
    db.db_path.rename(moved)
    moved.unlink()


@pytest.mark.parametrize('encoding', ['utf-8-sig', 'cp1251', 'utf-16'])
def test_txt_encodings_and_original_bytes(tmp_path, encoding):
    raw = 'Учебный пример. Ёжик и библиотека.'.encode(encoding)
    doc = extract_document('../../library.TXT', raw)
    assert doc['name'] == 'library.TXT'
    assert doc['pages'][0]['text'] == 'Учебный пример. Ёжик и библиотека.'
    db = DatabaseManager(tmp_path / 'db')
    db.initialize()
    db.add_document(doc)
    assert db.get_original(doc['id']) == raw


def test_docx_tables_keep_body_order():
    doc = Document()
    doc.add_paragraph('До таблицы')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Содержимое таблицы'
    doc.add_paragraph('После таблицы')
    stream = io.BytesIO()
    doc.save(stream)
    result = extract_document('test.docx', stream.getvalue())
    assert result['pages'][0]['text'] == 'До таблицы\nСодержимое таблицы\nПосле таблицы'
    assert result['pages'][0]['page'] is None


def test_pdf_page_numbers_are_not_renumbered_when_blank_pages_skipped(encoder):
    doc = extract_document('source.pdf', text_pdf(['First page', '', 'Third page']))
    assert [p['page'] for p in doc['pages']] == [1, 3]
    chunks = chunk_documents([doc], encoder.tokenizer, ChunkConfig())
    assert [c['page'] for c in chunks] == [1, 3]
    assert chunks[1]['document_name'] == 'source.pdf'
    assert '2' in doc['metadata']['warnings'][0]


@pytest.mark.parametrize(('name', 'raw', 'message'), [
    ('empty.txt', b'', 'пуст'), ('blank.txt', b' \n\t', 'не содержит текста'),
    ('broken.docx', b'not a zip', 'Не удалось прочитать'),
    ('broken.pdf', b'not a pdf', 'Не удалось прочитать'),
    ('image.png', b'not text', 'Поддерживаются'), ('binary.txt', b'\x00\x01', 'двоичные'),
])
def test_bad_files_are_explained(name, raw, message):
    with pytest.raises(DocumentError, match=message):
        extract_document(name, raw)


def test_scanned_and_encrypted_pdf_have_clear_errors():
    with pytest.raises(DocumentError, match='OCR'):
        extract_document('scan.pdf', text_pdf(['']))
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt('password')
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(DocumentError, match='паролем'):
        extract_document('locked.pdf', output.getvalue())


def test_chunks_have_overlap_original_spans_and_stable_ids(encoder):
    text = ' '.join(f'Слово{i}' for i in range(100))
    doc = extract_document('text.txt', text.encode())
    chunks = chunk_documents([doc], encoder.tokenizer, ChunkConfig(32, 8))
    assert chunks == chunk_documents([doc], encoder.tokenizer, ChunkConfig(32, 8))
    assert chunks[0]['text'].split()[-8:] == chunks[1]['text'].split()[:8]
    assert chunks[-1]['text'].endswith('Слово99')
    assert all(text[c['char_start']:c['char_end']].strip() == c['text'] for c in chunks)
    assert corpus_fingerprint([doc]) != corpus_fingerprint([{**doc, 'name': 'changed.txt'}])
